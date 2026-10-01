"""
Template generation utility for Registrar dataset uploads.
Generates downloadable Excel (.xlsx) templates with:
- Sheet 1: Header row & formatted example row
- Sheet 2: 'Notes' sheet explaining schema requirements, data types, constraints, and valid codes.
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
EXAMPLE_FONT = Font(name="Calibri", size=10, italic=True, color="334155")
NOTES_HEADER_FILL = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


TEMPLATES_CONFIG = {
    "courses": {
        "title": "Courses",
        "filename": "courses_template.xlsx",
        "headers": ["course_code", "title", "department_code", "program_code", "credits", "semester"],
        "example": ["STA308", "Statistical Data Cleaning & Live Dashboards", "STAT", "BS-STAT", 3, "Fall 2026"],
        "notes": [
            ("Upload Order", "Upload #1: Courses must be uploaded first (or configured by Admin) before Grades/Attendance."),
            ("course_code", "Unique alphanumeric course code (e.g. 'STA308', 'CS101'). Required."),
            ("title", "Full descriptive course title (e.g. 'Data Quality Profiling'). Required."),
            ("department_code", "Department code matching an existing Department (e.g. 'STAT', 'CS'). Required."),
            ("program_code", "Program code matching an existing Program (e.g. 'BS-STAT'). Optional."),
            ("credits", "Credit units (integer, typically 1 to 6). Default is 3."),
            ("semester", "Academic term/semester (e.g. 'Fall 2026', 'Spring 2026'). Required."),
        ],
    },
    "students": {
        "title": "Students",
        "filename": "students_template.xlsx",
        "headers": ["student_id", "first_name", "last_name", "email", "department_code", "program_code", "cohort", "status"],
        "example": ["STU1001", "Alex", "Taylor", "alex.taylor@university.edu", "STAT", "BS-STAT", "2024", "Active"],
        "notes": [
            ("Upload Order", "Upload #2: Students must be uploaded before uploading Grades or Attendance."),
            ("student_id", "Unique student identification string (e.g. 'STU1001'). Required."),
            ("first_name", "Student given name. Required."),
            ("last_name", "Student family name. Required."),
            ("email", "Valid institutional email address. Required and must be unique."),
            ("department_code", "Department code the student belongs to (e.g. 'STAT'). Required."),
            ("program_code", "Academic program code (e.g. 'BS-STAT'). Optional."),
            ("cohort", "Enrollment year/cohort (e.g. '2024')."),
            ("status", "Enrollment status: 'Active', 'Graduated', or 'Suspended'. Default is 'Active'."),
        ],
    },
    "grades": {
        "title": "Grades",
        "filename": "grades_template.xlsx",
        "headers": ["student_id", "course_code", "semester", "numerical_score"],
        "example": ["STU1001", "STA308", "Fall 2026", 88.5],
        "notes": [
            ("Upload Order", "Upload #3: Both Courses and Students must already exist in the database!"),
            ("student_id", "Must match an existing student in the portal. Required."),
            ("course_code", "Must match an existing course code in the portal. Required."),
            ("semester", "Academic term (e.g. 'Fall 2026'). Required."),
            ("numerical_score", "Percentage score between 0.0 and 100.0. Required."),
            ("Grading Scale", "Letter Grade & GPA points are auto-calculated (A>=90: 4.0, B+>=85: 3.5, B>=80: 3.0, C+>=75: 2.5, C>=70: 2.0, D>=60: 1.0, F<60: 0.0)."),
        ],
    },
    "attendance": {
        "title": "Attendance",
        "filename": "attendance_template.xlsx",
        "headers": ["student_id", "course_code", "date", "status", "session_topic"],
        "example": ["STU1001", "STA308", "2026-09-15", "Present", "Data Quality Profiling"],
        "notes": [
            ("Upload Order", "Upload #4: Both Courses and Students must already exist in the database!"),
            ("student_id", "Must match an existing student. Required."),
            ("course_code", "Must match an existing course. Required."),
            ("date", "Session date formatted as YYYY-MM-DD. Required."),
            ("status", "Valid status values: 'Present', 'Absent', 'Late', or 'Excused'. Required."),
            ("session_topic", "Brief lesson or lecture topic. Optional."),
        ],
    },
}


def generate_dataset_template(dataset_type: str) -> io.BytesIO:
    """Generate an in-memory .xlsx workbook for the requested dataset type."""
    if dataset_type not in TEMPLATES_CONFIG:
        raise ValueError(f"Unknown dataset template type: {dataset_type}")

    config = TEMPLATES_CONFIG[dataset_type]
    wb = openpyxl.Workbook()

    # Sheet 1: Data entry sheet
    ws_data = wb.active
    ws_data.title = config["title"]
    ws_data.views.sheetView[0].showGridLines = True

    # Write Headers
    headers = config["headers"]
    ws_data.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_data.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN

    # Write Example Row
    example = config["example"]
    ws_data.append(example)
    for col_idx in range(1, len(example) + 1):
        cell = ws_data.cell(row=2, column=col_idx)
        cell.font = EXAMPLE_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER_THIN

    # Auto-fit column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # Sheet 2: Notes sheet
    ws_notes = wb.create_sheet(title="Notes & Rules")
    ws_notes.views.sheetView[0].showGridLines = True
    ws_notes.append(["Field / Rule", "Description & Constraints"])
    for col_idx in (1, 2):
        cell = ws_notes.cell(row=1, column=col_idx)
        cell.fill = NOTES_HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item, desc in config["notes"]:
        ws_notes.append([item, desc])
        row_idx = ws_notes.max_row
        ws_notes.cell(row=row_idx, column=1).font = Font(name="Calibri", bold=True, color="0F172A")
        ws_notes.cell(row=row_idx, column=2).font = Font(name="Calibri", color="334155")
        ws_notes.cell(row=row_idx, column=1).border = BORDER_THIN
        ws_notes.cell(row=row_idx, column=2).border = BORDER_THIN

    ws_notes.column_dimensions["A"].width = 22
    ws_notes.column_dimensions["B"].width = 75

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
