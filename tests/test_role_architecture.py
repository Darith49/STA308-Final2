"""
Comprehensive Test Suite for Multi-Role University Portal Architecture:
- Role Checking & 403 Page Enforcement
- Server-Side Scoping (scope_for)
- Registrar Uploads, Schema Validation, Templates, and Error Files
- Atomic Cascading Rollbacks on Upload Deletion
- Scoped CSV Export with Formula Injection Protection
- Student Portal Data Isolation
- Audit Logging
"""

import io
import openpyxl
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError

from apps.accounts.models import UserProfile, UserRole, AuditLog
from apps.academic.models import Department, Program, Course, Student, Grade, Attendance
from apps.academic.scoping import scope_for, user_has_role
from apps.academic.importers import import_domain_dataset
from apps.uploads.models import Upload, UploadStatus, DatasetType, UploadError
from apps.uploads.validators import (
    validate_dataset_schema,
    validate_duplicate_hash,
    compute_file_sha256,
    sanitize_formula_injection,
)
from apps.uploads.templates_generator import generate_dataset_template


class MultiRolePortalTests(TestCase):

    def setUp(self):
        # 1. Departments & Programs
        self.dept_stat = Department.objects.create(code="STAT", name="Department of Statistics")
        self.dept_cs = Department.objects.create(code="CS", name="Department of Computer Science")

        self.prog_stat = Program.objects.create(department=self.dept_stat, code="BS-STAT", name="BSc Statistics")
        self.prog_cs = Program.objects.create(department=self.dept_cs, code="BS-CS", name="BSc Computer Science")

        # 2. Courses
        self.course_stat = Course.objects.create(
            department=self.dept_stat,
            program=self.prog_stat,
            code="STA308",
            title="Statistical Data Cleaning",
            credits=3,
            semester="Fall 2026",
        )
        self.course_cs = Course.objects.create(
            department=self.dept_cs,
            program=self.prog_cs,
            code="CS101",
            title="Algorithms",
            credits=4,
            semester="Fall 2026",
        )

        # 3. Users for all 4 roles
        # Admin
        self.admin_user = User.objects.create_user(username="admin_user", password="password123", is_staff=True)
        self.admin_user.profile.role = UserRole.ADMIN
        self.admin_user.profile.save()

        # Registrar
        self.registrar_user = User.objects.create_user(username="reg_user", password="password123")
        self.registrar_user.profile.role = UserRole.REGISTRAR
        self.registrar_user.profile.save()

        # Department Head (STAT)
        self.dept_head = User.objects.create_user(username="dept_head", password="password123")
        self.dept_head.profile.role = UserRole.DEPT_HEAD
        self.dept_head.profile.department = "Department of Statistics"
        self.dept_head.profile.save()
        self.dept_stat.head = self.dept_head
        self.dept_stat.save()

        # Student User & Record
        self.student_user = User.objects.create_user(username="student_user", email="student@university.edu", password="password123")
        self.student_user.profile.role = UserRole.STUDENT
        self.student_user.profile.save()

        self.student_record = Student.objects.create(
            student_id="STU1001",
            user=self.student_user,
            first_name="Jane",
            last_name="Doe",
            email="student@university.edu",
            department=self.dept_stat,
            program=self.prog_stat,
            cohort="2024",
            status="Active",
        )

        # Other Student (CS)
        self.student_cs = Student.objects.create(
            student_id="STU2001",
            first_name="Bob",
            last_name="Smith",
            email="bob@cs.edu",
            department=self.dept_cs,
            program=self.prog_cs,
            cohort="2024",
            status="Active",
        )

        # Grades
        self.grade_stat = Grade.objects.create(
            student=self.student_record,
            course=self.course_stat,
            semester="Fall 2026",
            numerical_score=92.0,
        )
        self.grade_cs = Grade.objects.create(
            student=self.student_cs,
            course=self.course_cs,
            semester="Fall 2026",
            numerical_score=75.0,
        )

        # Clients
        self.client_anon = Client()
        self.client_admin = Client()
        self.client_admin.login(username="admin_user", password="password123")

        self.client_reg = Client()
        self.client_reg.login(username="reg_user", password="password123")

        self.client_head = Client()
        self.client_head.login(username="dept_head", password="password123")

        self.client_student = Client()
        self.client_student.login(username="student_user", password="password123")

    # --------------------------------------------------------------------------
    # 1. Role Authorization & 403 Enforcement
    # --------------------------------------------------------------------------

    def test_anonymous_redirects_to_login(self):
        """Unauthenticated requests are directed to the login page."""
        protected_urls = [
            "/academic/admin/overview/",
            "/academic/department/dashboard/",
            "/academic/student/portal/",
            "/uploads/new/",
        ]
        for url in protected_urls:
            res = self.client_anon.get(url)
            self.assertEqual(res.status_code, 302, f"Failed for {url}")
            self.assertIn("/accounts/login/", res.url)

    def test_role_permissions_admin_dashboard(self):
        """Only Admin can access /academic/admin/overview/. Others receive HTTP 403."""
        # Admin succeeds
        res_admin = self.client_admin.get("/academic/admin/overview/")
        self.assertEqual(res_admin.status_code, 200)

        # Dept Head gets 403
        res_head = self.client_head.get("/academic/admin/overview/")
        self.assertEqual(res_head.status_code, 403)

        # Student gets 403
        res_student = self.client_student.get("/academic/admin/overview/")
        self.assertEqual(res_student.status_code, 403)

    def test_role_home_redirect_dispatch(self):
        """Root URL '/' dispatches each authenticated user to their role landing page."""
        # Admin -> Admin dashboard
        res = self.client_admin.get("/", follow=False)
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/academic/admin/overview/")

        # Registrar -> Uploads list
        res = self.client_reg.get("/", follow=False)
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/uploads/")

        # Dept Head -> Department dashboard
        res = self.client_head.get("/", follow=False)
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/academic/department/dashboard/")

        # Student -> Student portal
        res = self.client_student.get("/", follow=False)
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/academic/student/portal/")

    # --------------------------------------------------------------------------
    # 2. Server-Side Data Scoping (scope_for)
    # --------------------------------------------------------------------------

    def test_department_head_server_side_scoping(self):
        """Dept Head for STAT only sees STAT records, never CS records."""
        # Scoped student queryset
        scoped_students = scope_for(self.dept_head, Student)
        self.assertIn(self.student_record, scoped_students)
        self.assertNotIn(self.student_cs, scoped_students)

        # Scoped course queryset
        scoped_courses = scope_for(self.dept_head, Course)
        self.assertIn(self.course_stat, scoped_courses)
        self.assertNotIn(self.course_cs, scoped_courses)

    def test_student_data_isolation(self):
        """Student only sees their own student record and grades."""
        # Student scoped
        scoped_students = scope_for(self.student_user, Student)
        self.assertEqual(scoped_students.count(), 1)
        self.assertEqual(scoped_students.first(), self.student_record)

        # Student grades scoped
        scoped_grades = scope_for(self.student_user, Grade)
        self.assertIn(self.grade_stat, scoped_grades)
        self.assertNotIn(self.grade_cs, scoped_grades)

    # --------------------------------------------------------------------------
    # 3. Registrar Template Downloads
    # --------------------------------------------------------------------------

    def test_registrar_template_downloads(self):
        """Registrar can download valid .xlsx templates for all 4 dataset types."""
        types = ["courses", "students", "grades", "attendance"]
        for t in types:
            res = self.client_reg.get(f"/uploads/templates/{t}/")
            self.assertEqual(res.status_code, 200)
            self.assertEqual(
                res["Content-Type"],
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            # Verify valid openpyxl readable workbook
            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            self.assertIn("Notes & Rules", wb.sheetnames)

    # --------------------------------------------------------------------------
    # 4. Schema Validation & Duplicate Hash Checks
    # --------------------------------------------------------------------------

    def test_schema_validation_rejects_missing_columns(self):
        """Uploaded file missing required columns is rejected before saving."""
        # Create an in-memory workbook with invalid headers
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["wrong_column_1", "wrong_column_2"])
        ws.append(["val1", "val2"])
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)

        fake_file = SimpleUploadedFile("bad_courses.xlsx", out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with self.assertRaises(ValidationError) as ctx:
            validate_dataset_schema(fake_file, "courses")
        self.assertIn("Missing required column", str(ctx.exception))

    def test_duplicate_file_hash_prevention(self):
        """Duplicate file hashes are rejected."""
        dummy_hash = "abcdef0123456789"
        Upload.objects.create(
            owner=self.registrar_user,
            original_filename="first_upload.xlsx",
            sha256=dummy_hash,
            status=UploadStatus.DONE,
        )
        with self.assertRaises(ValidationError):
            validate_duplicate_hash(dummy_hash)

    # --------------------------------------------------------------------------
    # 5. Cascading Rollbacks on Upload Deletion
    # --------------------------------------------------------------------------

    def test_upload_deletion_cascades_and_rolls_back_domain_rows(self):
        """Deleting an upload removes the upload and all rows linked to it."""
        upload = Upload.objects.create(
            owner=self.registrar_user,
            dataset_type=DatasetType.COURSES,
            original_filename="courses_batch.xlsx",
            status=UploadStatus.DONE,
        )
        # Create a course linked to this upload
        imported_course = Course.objects.create(
            department=self.dept_stat,
            code="STA499",
            title="Senior Capstone",
            credits=3,
            semester="Fall 2026",
            upload=upload,
        )
        # Create an error record
        err = UploadError.objects.create(
            upload=upload,
            row_number=5,
            column_name="code",
            raw_value="BAD",
            reason="Invalid code",
        )

        self.assertTrue(Course.objects.filter(id=imported_course.id).exists())
        self.assertTrue(UploadError.objects.filter(id=err.id).exists())

        # Registrar deletes their own upload
        res = self.client_reg.post(f"/uploads/{upload.id}/delete/")
        self.assertEqual(res.status_code, 302)

        # Verify domain course and error were rolled back!
        self.assertFalse(Course.objects.filter(id=imported_course.id).exists())
        self.assertFalse(UploadError.objects.filter(id=err.id).exists())
        self.assertFalse(Upload.objects.filter(id=upload.id).exists())

    # --------------------------------------------------------------------------
    # 6. CSV Export Formula Injection Protection
    # --------------------------------------------------------------------------

    def test_csv_export_sanitizes_formula_injection(self):
        """Formula injection strings in exported tables are neutralized with leading quote."""
        # Inject dangerous student name
        dangerous_student = Student.objects.create(
            student_id="STU9999",
            first_name="=cmd|' /C calc'!A0",
            last_name="+SUM(1,2)",
            email="hack@edu.com",
            department=self.dept_stat,
        )
        res = self.client_head.get("/academic/department/export/students/")
        self.assertEqual(res.status_code, 200)
        content = res.content.decode("utf-8")

        # Verify sanitized with leading single-quote
        self.assertIn("'=cmd|' /C calc'!A0", content)
        self.assertIn("'+SUM(1,2)", content)

    # --------------------------------------------------------------------------
    # 7. Audit Logging
    # --------------------------------------------------------------------------

    def test_audit_log_created_on_user_management(self):
        """Admin creating a user logs an entry in AuditLog."""
        initial_count = AuditLog.objects.count()
        res = self.client_admin.post(
            "/academic/admin/users/create/",
            {
                "username": "new_lecturer",
                "email": "lecturer@university.edu",
                "first_name": "Alan",
                "last_name": "Turing",
                "password": "secure_password_123",
                "role": UserRole.DEPT_HEAD,
                "department": self.dept_stat.id,
            },
        )
        self.assertEqual(res.status_code, 302)
        self.assertTrue(User.objects.filter(username="new_lecturer").exists())
        self.assertEqual(AuditLog.objects.count(), initial_count + 1)
        latest_log = AuditLog.objects.first()
        self.assertEqual(latest_log.action, "USER_CREATE")
        self.assertEqual(latest_log.user, self.admin_user)
