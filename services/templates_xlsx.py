import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEMPLATE_DEFINITIONS = {
    'students': {
        'title': 'Students Template',
        'columns': ['student_id', 'name', 'gender', 'department', 'program', 'year', 'enrollment_date', 'email'],
        'example_rows': [
            ['STU1001', 'Alice Johnson', 'Female', 'CS', 'Bachelor of Computer Science', 1, '2023-09-01', 'alice.j@university.edu'],
            ['STU1002', 'Bob Smith', 'Male', 'ENG', 'Mechanical Engineering', 2, '2022-09-01', 'bob.smith@university.edu'],
        ],
        'notes': [
            ('student_id', 'Text / Alphanumeric', 'Required. Unique student identifier (e.g. STU1001). Rows without ID are rejected.'),
            ('name', 'Text', 'Optional / Recommended. Student full name. Will be title-cased.'),
            ('gender', 'Text', 'Optional. Male, Female, Other. Standardized automatically (m/f/male/female).'),
            ('department', 'Text / Code', 'Optional. Department code or name (e.g. CS, ENG, BUS).'),
            ('program', 'Text', 'Optional. Degree program name.'),
            ('year', 'Integer (1-6)', 'Optional. Current year of study (1 to 6). Invalid numbers logged as warnings and set to null.'),
            ('enrollment_date', 'Date (YYYY-MM-DD)', 'Optional. Must not be in the future. Formats like YYYY-MM-DD or DD/MM/YYYY supported.'),
            ('email', 'Email Text', 'Optional. Malformed emails logged as warnings and set to null.'),
        ]
    },
    'courses': {
        'title': 'Courses Template',
        'columns': ['course_code', 'title', 'credits', 'department'],
        'example_rows': [
            ['CS101', 'Introduction to Computer Science', 3, 'CS'],
            ['ENG201', 'Thermodynamics I', 4, 'ENG'],
            ['BUS105', 'Principles of Management', 3, 'BUS'],
        ],
        'notes': [
            ('course_code', 'Text', 'Required. Unique course code (e.g. CS101). Uppercased automatically.'),
            ('title', 'Text', 'Optional. Title of the course. Will be title-cased.'),
            ('credits', 'Integer (1-12)', 'Optional. Credit hours (default is 3).'),
            ('department', 'Text', 'Department code or faculty name.'),
        ]
    },
    'grades': {
        'title': 'Grades Template',
        'columns': ['student_id', 'course_code', 'semester', 'score'],
        'example_rows': [
            ['STU1001', 'CS101', '2024-Fall', 88.5],
            ['STU1002', 'CS101', '2024-Fall', 74.0],
            ['STU1001', 'ENG201', '2024-Fall', 92.0],
        ],
        'notes': [
            ('student_id', 'Text', 'Required. Must correspond to a registered student. Unknown IDs rejected.'),
            ('course_code', 'Text', 'Required. Must correspond to an existing course. Unknown codes rejected.'),
            ('semester', 'Text', 'Optional. Academic semester (e.g. 2024-Fall, 2024-Spring). Defaults to 2024-Fall.'),
            ('score', 'Decimal (0-100)', 'Required. Numeric grade between 0.0 and 100.0. Scores outside range rejected. Letter grades are calculated automatically.'),
        ]
    },
    'attendance': {
        'title': 'Attendance Template',
        'columns': ['student_id', 'course_code', 'date', 'status'],
        'example_rows': [
            ['STU1001', 'CS101', '2024-10-15', 'P'],
            ['STU1002', 'CS101', '2024-10-15', 'A'],
            ['STU1001', 'CS101', '2024-10-17', 'L'],
        ],
        'notes': [
            ('student_id', 'Text', 'Required. Must correspond to an existing student.'),
            ('course_code', 'Text', 'Required. Must correspond to an existing course.'),
            ('date', 'Date (YYYY-MM-DD)', 'Required. Class session date. Future dates are rejected.'),
            ('status', 'Text / Code', 'Required. Allowed values: Present (P, 1), Absent (A, 0), Late (L).'),
        ]
    }
}

def generate_template_xlsx(dataset_type):
    """
    Generates a professionally styled openpyxl Workbook for dataset_type.
    Returns bytes buffer of XLSX file.
    """
    if dataset_type not in TEMPLATE_DEFINITIONS:
        raise ValueError(f"Unknown dataset type '{dataset_type}'")

    defn = TEMPLATE_DEFINITIONS[dataset_type]
    wb = openpyxl.Workbook()

    # Sheet 1: Data
    ws_data = wb.active
    ws_data.title = "Data"

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = Font(name='Segoe UI', size=11, bold=True)
    navy_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    light_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # Headers
    for col_idx, col_name in enumerate(defn['columns'], start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Example rows
    for row_idx, row_data in enumerate(defn['example_rows'], start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10)
            cell.fill = light_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    # Auto-fit column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 5, 16)

    # Sheet 2: Notes
    ws_notes = wb.create_sheet(title="Notes")
    ws_notes.cell(row=1, column=1, value="Column").font = header_font
    ws_notes.cell(row=1, column=1).fill = navy_fill
    ws_notes.cell(row=1, column=2, value="Data Type").font = header_font
    ws_notes.cell(row=1, column=2).fill = navy_fill
    ws_notes.cell(row=1, column=3, value="Description & Rules").font = header_font
    ws_notes.cell(row=1, column=3).fill = navy_fill

    for r_idx, (col_name, dtype, note_text) in enumerate(defn['notes'], start=2):
        c1 = ws_notes.cell(row=r_idx, column=1, value=col_name)
        c2 = ws_notes.cell(row=r_idx, column=2, value=dtype)
        c3 = ws_notes.cell(row=r_idx, column=3, value=note_text)
        for c in (c1, c2, c3):
            c.font = Font(name='Segoe UI', size=10)
            c.border = thin_border
            c.alignment = Alignment(vertical='top')

    ws_notes.column_dimensions['A'].width = 20
    ws_notes.column_dimensions['B'].width = 22
    ws_notes.column_dimensions['C'].width = 65

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
