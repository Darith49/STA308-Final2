import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from django.db import transaction
from django.utils import timezone

from uploads.models import Upload, UploadError, UploadStatus, AuditLog
from academics.models import Department, Program, Course, Student, Enrollment, Attendance
from .cleaners import StudentsCleaner, CoursesCleaner, GradesCleaner, AttendanceCleaner

def run_import_pipeline(upload_id):
    """
    Executes the complete cleaning, validation, foreign key resolution,
    and bulk import pipeline for a given Upload instance.
    Can be run synchronously or via background task.
    """
    try:
        upload = Upload.objects.get(id=upload_id)
    except Upload.DoesNotExist:
        return None

    upload.status = UploadStatus.PROCESSING
    upload.progress_percent = 10
    upload.save(update_fields=['status', 'progress_percent'])

    try:
        # Step 1: Read Excel file
        file_path = upload.file.path
        df = pd.read_excel(file_path, sheet_name=0, engine='openpyxl')
        upload.progress_percent = 25
        upload.save(update_fields=['progress_percent'])

        cleaner_cls_map = {
            'students': StudentsCleaner,
            'courses': CoursesCleaner,
            'grades': GradesCleaner,
            'attendance': AttendanceCleaner,
        }

        cleaner_cls = cleaner_cls_map.get(upload.dataset_type)
        if not cleaner_cls:
            raise ValueError(f"Unknown dataset type '{upload.dataset_type}'")

        cleaner = cleaner_cls()
        clean_df, errors, report = cleaner.clean(df)
        upload.progress_percent = 50
        upload.save(update_fields=['progress_percent'])

        # Step 2: Database FK Resolution and Validation
        objects_to_create = []
        extra_errors = []

        if upload.dataset_type == 'students':
            # Pre-cache or create departments and programs
            dept_cache = {d.code.upper(): d for d in Department.objects.all()}
            prog_cache = {(p.department.code.upper(), p.name.lower()): p for p in Program.objects.select_related('department')}

            for idx, row in clean_df.iterrows():
                dept_code = str(row['department']).strip().upper()
                if dept_code not in dept_cache:
                    new_dept, _ = Department.objects.get_or_create(code=dept_code, defaults={'name': dept_code})
                    dept_cache[dept_code] = new_dept
                dept_obj = dept_cache[dept_code]

                prog_name = str(row['program']).strip()
                prog_key = (dept_code, prog_name.lower())
                if prog_key not in prog_cache:
                    new_prog, _ = Program.objects.get_or_create(department=dept_obj, name=prog_name)
                    prog_cache[prog_key] = new_prog
                prog_obj = prog_cache[prog_key]

                # Extract sanitized values (avoiding NaN in integer/date fields)
                year_val = row.get('year')
                cleaned_year = int(float(year_val)) if (year_val is not None and pd.notna(year_val)) else None
                enrol_val = row.get('enrollment_date') if pd.notna(row.get('enrollment_date')) else None
                email_raw = row.get('email')
                email_val = str(email_raw).strip() if (email_raw is not None and pd.notna(email_raw) and str(email_raw).lower() not in ('nan', 'none')) else None

                # Update existing student or stage new student
                sid = row['student_id']
                existing_stu = Student.objects.filter(student_id=sid).first()
                if existing_stu:
                    # Update student record to link to current upload
                    existing_stu.name = row['name']
                    existing_stu.gender = row['gender']
                    existing_stu.program = prog_obj
                    existing_stu.year = cleaned_year
                    existing_stu.enrollment_date = enrol_val
                    existing_stu.email = email_val
                    existing_stu.upload = upload
                    existing_stu.save()
                else:
                    objects_to_create.append(
                        Student(
                            student_id=sid,
                            name=row['name'],
                            gender=row['gender'],
                            program=prog_obj,
                            year=cleaned_year,
                            enrollment_date=enrol_val,
                            email=email_val,
                            upload=upload
                        )
                    )

        elif upload.dataset_type == 'courses':
            dept_cache = {d.code.upper(): d for d in Department.objects.all()}

            for idx, row in clean_df.iterrows():
                dept_code = str(row['department']).strip().upper()
                if dept_code not in dept_cache:
                    new_dept, _ = Department.objects.get_or_create(code=dept_code, defaults={'name': dept_code})
                    dept_cache[dept_code] = new_dept
                dept_obj = dept_cache[dept_code]

                code = row['course_code']
                credits_val = int(row['credits']) if pd.notna(row.get('credits')) else 3
                existing_course = Course.objects.filter(code=code).first()
                if existing_course:
                    existing_course.title = row['title']
                    existing_course.credits = credits_val
                    existing_course.department = dept_obj
                    existing_course.upload = upload
                    existing_course.save()
                else:
                    objects_to_create.append(
                        Course(
                            code=code,
                            title=row['title'],
                            credits=credits_val,
                            department=dept_obj,
                            upload=upload
                        )
                    )

        elif upload.dataset_type == 'grades':
            # Pre-load known students and courses
            all_students = {s.student_id: s for s in Student.objects.all()}
            all_courses = {c.code.upper(): c for c in Course.objects.all()}
            report.setdefault('unknown_students', 0)
            report.setdefault('unknown_courses', 0)

            valid_grades = []
            for idx, row in clean_df.iterrows():
                row_num = idx + 2
                sid = row['student_id']
                ccode = row['course_code'].upper()

                if sid not in all_students:
                    extra_errors.append({
                        'row_number': row_num,
                        'column': 'student_id',
                        'value': sid,
                        'message': f"Student ID '{sid}' not found in database. Import students before grades."
                    })
                    report['unknown_students'] += 1
                    report['rows_rejected'] += 1
                    continue

                if ccode not in all_courses:
                    extra_errors.append({
                        'row_number': row_num,
                        'column': 'course_code',
                        'value': ccode,
                        'message': f"Course code '{ccode}' not found in database. Import courses before grades."
                    })
                    report['unknown_courses'] += 1
                    report['rows_rejected'] += 1
                    continue

                student_obj = all_students[sid]
                course_obj = all_courses[ccode]
                semester_val = row['semester']

                # Check if existing enrollment exists
                existing_enr = Enrollment.objects.filter(student=student_obj, course=course_obj, semester=semester_val).first()
                if existing_enr:
                    existing_enr.score = row['score']
                    existing_enr.grade_letter = row['grade_letter']
                    existing_enr.upload = upload
                    existing_enr.save()
                else:
                    objects_to_create.append(
                        Enrollment(
                            student=student_obj,
                            course=course_obj,
                            semester=semester_val,
                            score=row['score'],
                            grade_letter=row['grade_letter'],
                            upload=upload
                        )
                    )

        elif upload.dataset_type == 'attendance':
            all_students = {s.student_id: s for s in Student.objects.all()}
            all_courses = {c.code.upper(): c for c in Course.objects.all()}
            report.setdefault('unknown_students', 0)
            report.setdefault('unknown_courses', 0)

            for idx, row in clean_df.iterrows():
                row_num = idx + 2
                sid = row['student_id']
                ccode = row['course_code'].upper()

                if sid not in all_students:
                    extra_errors.append({
                        'row_number': row_num,
                        'column': 'student_id',
                        'value': sid,
                        'message': f"Student ID '{sid}' not found in database. Import students before attendance."
                    })
                    report['unknown_students'] += 1
                    report['rows_rejected'] += 1
                    continue

                if ccode not in all_courses:
                    extra_errors.append({
                        'row_number': row_num,
                        'column': 'course_code',
                        'value': ccode,
                        'message': f"Course code '{ccode}' not found in database. Import courses before attendance."
                    })
                    report['unknown_courses'] += 1
                    report['rows_rejected'] += 1
                    continue

                objects_to_create.append(
                    Attendance(
                        student=all_students[sid],
                        course=all_courses[ccode],
                        date=row['date'],
                        status=row['status'],
                        upload=upload
                    )
                )

        upload.progress_percent = 75
        upload.save(update_fields=['progress_percent'])

        # Step 3: Atomic Bulk Insert & Error Logging
        all_errors = errors + extra_errors
        rows_ok_count = len(clean_df) - len(extra_errors)
        report['rows_ok'] = max(rows_ok_count, 0)

        with transaction.atomic():
            if objects_to_create:
                model_cls = objects_to_create[0].__class__
                model_cls.objects.bulk_create(objects_to_create, batch_size=500)

            error_objs = [
                UploadError(
                    upload=upload,
                    row_number=err['row_number'],
                    column=err.get('column', ''),
                    value=err.get('value', ''),
                    message=err['message']
                )
                for err in all_errors
            ]
            if error_objs:
                UploadError.objects.bulk_create(error_objs, batch_size=500)

            # Update upload record
            upload.status = UploadStatus.DONE
            upload.rows_in = report.get('rows_in', 0)
            upload.rows_ok = report.get('rows_ok', 0)
            upload.rows_rejected = report.get('rows_rejected', 0)
            upload.report = report
            upload.progress_percent = 100
            upload.save()

            # Record audit log
            AuditLog.objects.create(
                user=upload.uploaded_by,
                action='UPLOAD_PROCESSED',
                target=f"Upload #{upload.id} ({upload.get_dataset_type_display()})",
                details={
                    'dataset_type': upload.dataset_type,
                    'rows_in': upload.rows_in,
                    'rows_ok': upload.rows_ok,
                    'rows_rejected': upload.rows_rejected,
                }
            )

        return upload

    except Exception as e:
        upload.status = UploadStatus.FAILED
        upload.error_message = str(e)
        upload.save(update_fields=['status', 'error_message'])
        return upload


def generate_error_xlsx(upload):
    """
    Generates a downloadable XLSX containing all rejected/warning rows with error reasons.
    """
    errors = UploadError.objects.filter(upload=upload).order_by('row_number')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rejected Rows & Warnings"

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    red_fill = PatternFill(start_color='DC2626', end_color='DC2626', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    headers = ["Row Number", "Column", "Value Given", "Error / Rejection Reason"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = red_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r_idx, err in enumerate(errors, start=2):
        c1 = ws.cell(row=r_idx, column=1, value=err.row_number)
        c2 = ws.cell(row=r_idx, column=2, value=err.column)
        c3 = ws.cell(row=r_idx, column=3, value=err.value)
        c4 = ws.cell(row=r_idx, column=4, value=err.message)
        for c in (c1, c2, c3, c4):
            c.font = Font(name='Segoe UI', size=10)
            c.border = thin_border

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
